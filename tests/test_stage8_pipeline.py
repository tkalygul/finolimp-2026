import importlib.util
import json
from pathlib import Path
import tempfile
import unittest
from openpyxl import load_workbook
import pandas as pd
from reconcile import run_pipeline,check_inputs
from src.schema import ValidationError
from src.load_registry import InputFileError
from verify_run import compare,verify


class PipelineTests(unittest.TestCase):
    def inputs(self,root):
        root.mkdir()
        pd.DataFrame([dict(folder='SkyWay Travel',period_start='2026-01-01',period_end='2026-01-31',
            act_status='',date='2026-01-15',doc='Реализация',invoice='СЧ-1',ticket_cell='2461885130',
            pax='TEST/NAME',debet=100,credit=0,saldo_start=0,saldo_end=100)]).to_csv(root/'acts.csv',index=False)
        pd.DataFrame([dict(date='2026-01-15 12:00:00',txn_id='1',amount=-100,amount_kgs=-100,
            balance_after=-100,agent='SkyWay Travel',agreement_id='AGR-1',kind='выкуп',
            tickets='6042461885130',currency='KGS',creator='etm-bot',comment='')]).to_csv(root/'etm.csv',index=False)
        pd.DataFrame([dict(date='15.01.2026',employee='test.employee',kind='продажа',party='SkyWay Travel',
            tickets='6042461885130',pax='TEST/NAME',pnr='ABC123',airline='AR',route='FRU-TAS',
            pay_cell='100 KGS',rate_usd=88.17,rate_eur=94.59,rate_rub=1,rate_kzt=.171)]).to_csv(root/'registry.csv',index=False)

    def test_missing_input_is_explained(self):
        with tempfile.TemporaryDirectory() as folder:
            with self.assertRaisesRegex(InputFileError,'acts.csv'): check_inputs(folder)

    def test_duplicate_registry_rejected(self):
        with tempfile.TemporaryDirectory() as folder:
            data=Path(folder)/'data'; self.inputs(data)
            (data/'registry_extra.csv').write_bytes((data/'registry.csv').read_bytes())
            with self.assertRaisesRegex(InputFileError,'несколько'): check_inputs(data)

    def test_missing_column_is_explained(self):
        with tempfile.TemporaryDirectory() as folder:
            data=Path(folder)/'data'; self.inputs(data)
            f=pd.read_csv(data/'acts.csv'); f.drop(columns=['saldo_end']).to_csv(data/'acts.csv',index=False)
            with self.assertRaisesRegex(ValidationError,'saldo_end'): check_inputs(data)

    @unittest.skipUnless(importlib.util.find_spec('pyarrow'),'Нужен pyarrow для полного запуска')
    def test_fresh_runs_repeatability_single_month_and_failed_run_preservation(self):
        with tempfile.TemporaryDirectory() as folder:
            data=Path(folder)/'data'; self.inputs(data); out=Path(folder)/'output'
            first=run_pipeline(data,out)
            second=run_pipeline(data,out)
            self.assertNotEqual(first,second)
            self.assertGreater(compare(first,second),10)
            self.assertEqual(verify(second)['ml']['status'],'unavailable')
            w=load_workbook(second/'reconciliation_report.xlsx',read_only=True)
            self.assertIn('Действия бухгалтера',w.sheetnames)
            self.assertIn('Качество модели',w.sheetnames); w.close()
            previous=(out/'reconciliation_report.xlsx').read_bytes()
            pointer=(out/'latest_run.json').read_bytes()
            f=pd.read_csv(data/'acts.csv'); f['debet']='ошибка'; f.to_csv(data/'acts.csv',index=False)
            with self.assertRaises(ValidationError): run_pipeline(data,out)
            self.assertEqual((out/'reconciliation_report.xlsx').read_bytes(),previous)
            self.assertEqual((out/'latest_run.json').read_bytes(),pointer)
            manifests=[json.loads(p.read_text(encoding='utf-8')) for p in (out/'runs').glob('*/run_manifest.json')]
            self.assertEqual(sum(m['status']=='failed' for m in manifests),1)

    def test_output_cannot_pollute_inputs(self):
        with tempfile.TemporaryDirectory() as folder:
            with self.assertRaisesRegex(InputFileError,'вне папки'): run_pipeline(folder,Path(folder)/'output')


if __name__=='__main__': unittest.main()
