import unittest
import json
import tempfile
from pathlib import Path
from contextlib import ExitStack
from unittest.mock import patch
import pandas as pd
from src.accountant_actions import build_accountant_actions
from src.schema import ValidationError


class AccountantTests(unittest.TestCase):
    def operation(self, **updates):
        row=dict(match_id='m1',subagent_id='a',period='2026-01',is_error=1,error_type='etm_wrong_amount',
                 error_owner='bot',confidence='medium',amount_1c=100.,amount_etm=80.,flags='[]',
                 doc_1c='["doc1"]',txn_ids='["txn1"]',reason='wrong amount')
        row.update(updates); return row

    def bridge(self, **updates):
        row=dict(subagent_id='a',period='2026-01',closing_difference=20.,opening_difference=0.,unexplained=0.)
        row.update(updates); return row

    def run_actions(self, row=None, bridge=None, anomalies=None):
        return build_accountant_actions(pd.DataFrame([row or self.operation()]),
                                        pd.DataFrame([bridge or self.bridge()]),anomalies)

    def test_etm_sign_and_scenario(self):
        summary,detail=self.run_actions()
        self.assertEqual(detail.proposed_delta.iloc[0],-20)
        self.assertEqual(summary.expected_difference_after_proposals.iloc[0],0)
        self.assertEqual(summary.approved_etm_delta.iloc[0],0)
        self.assertEqual(summary.remaining_difference_without_approval.iloc[0],20)

    def test_worsening_total_keeps_correct_local_proposal(self):
        s,d=self.run_actions(bridge=self.bridge(closing_difference=-100))
        self.assertEqual(d.proposed_delta.iloc[0],-20)
        self.assertEqual(s.expected_difference_after_proposals.iloc[0],-120)
        self.assertEqual(s.scenario_absolute_change.iloc[0],20)
        self.assertEqual(s.scenario_assessment.iloc[0],'Увеличивается')
        self.assertIn('истории проводок',s.scenario_warning.iloc[0])
        self.assertIn(s.scenario_warning.iloc[0],s.clarify.iloc[0])

    def test_improving_scenario_has_no_warning(self):
        s,_=self.run_actions()
        self.assertEqual(s.scenario_assessment.iloc[0],'Уменьшается')
        self.assertEqual(s.scenario_absolute_change.iloc[0],-20)
        self.assertEqual(s.scenario_warning.iloc[0],'')

    def test_unknown_scenario_is_not_assessed_as_improved(self):
        s,_=self.run_actions(bridge=self.bridge(closing_difference=float('nan')))
        self.assertEqual(s.scenario_assessment.iloc[0],'Неизвестно')
        self.assertTrue(pd.isna(s.scenario_absolute_change.iloc[0]))
        self.assertEqual(s.scenario_warning.iloc[0],'')

    def test_overcharged_etm_increases_deposit(self):
        summary,_=self.run_actions(self.operation(amount_etm=120),self.bridge(closing_difference=-20))
        self.assertEqual(summary.etm_delta_after_confirmation.iloc[0],20)
        self.assertEqual(summary.expected_difference_after_proposals.iloc[0],0)

    def test_missing_etm_uses_zero_movement(self):
        _,d=self.run_actions(self.operation(error_type='etm_not_executed',amount_etm=float('nan')))
        self.assertEqual(d.proposed_delta.iloc[0],-100)

    def test_onec_void_reduces_debt(self):
        s,d=self.run_actions(self.operation(error_type='1c_voided_posted',error_owner='1c',amount_etm=0),
                            self.bridge(closing_difference=100))
        self.assertEqual(d.target.iloc[0],'1С')
        self.assertEqual(s.onec_debt_delta_after_confirmation.iloc[0],-100)
        self.assertEqual(s.expected_difference_after_proposals.iloc[0],0)

    def test_duplicate_anomalies_do_not_add_money(self):
        a=pd.DataFrame([dict(anomaly_id='a1',match_ids='["m1"]',status='suspected',subagent_id='a',amount_kgs=100),
                        dict(anomaly_id='a2',match_ids='["m1"]',status='suspected',subagent_id='a',amount_kgs=100)])
        s,d=self.run_actions(anomalies=a)
        self.assertEqual(len(d),1)
        self.assertEqual(s.etm_delta_after_confirmation.iloc[0],0)
        self.assertEqual(s.anomaly_count.iloc[0],2)
        self.assertTrue(pd.isna(d.proposed_delta.iloc[0]))

    def test_period_difference_requires_review(self):
        _,d=self.run_actions(self.operation(flags='["period_difference"]'))
        self.assertEqual(d.target.iloc[0],'выяснить')

    def test_registry_error_never_changes_balance(self):
        s,d=self.run_actions(self.operation(error_type='registry_parse_error',error_owner='agent',confidence='high'))
        self.assertEqual(s.etm_delta_after_confirmation.iloc[0],0)
        self.assertIn('реестра',d.action.iloc[0])

    def test_unknown_balance_is_not_zero(self):
        s,_=self.run_actions(bridge=self.bridge(closing_difference=float('nan'),unexplained=float('nan')))
        self.assertTrue(pd.isna(s.expected_difference_after_proposals.iloc[0]))
        self.assertEqual(s.unknown_months.iloc[0],1)

    def test_one_line_per_agent_last_balance_not_sum(self):
        b=pd.DataFrame([self.bridge(),self.bridge(period='2026-02',closing_difference=30)])
        s,_=build_accountant_actions(pd.DataFrame([self.operation()]),b)
        self.assertEqual(len(s),1)
        self.assertEqual(s.closing_difference.iloc[0],30)
        self.assertEqual(s.expected_difference_after_proposals.iloc[0],10)

    def test_duplicate_groups_rejected(self):
        with self.assertRaises(ValidationError):
            build_accountant_actions(pd.DataFrame([self.operation()]*2),pd.DataFrame([self.bridge()]))

    def test_no_errors_still_has_summary(self):
        s,d=self.run_actions(self.operation(is_error=0,error_type='ok'),self.bridge(closing_difference=0))
        self.assertTrue(d.empty)
        self.assertEqual(s.status.iloc[0],'согласовано')

    def test_internal_balance_problem_is_explicit(self):
        s,_=self.run_actions(self.operation(is_error=0,error_type='ok'),
                            self.bridge(closing_difference=0,check_1c=10))
        self.assertIn('внутренний баланс акта',s.clarify.iloc[0])

    def test_report_integration_saves_actions_and_excel(self):
        from src.report import build_report
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as folder:
            root=Path(folder); (root/'p5').mkdir(); (root/'reconciliation').mkdir()
            pd.DataFrame([self.operation()]).to_csv(root/'p5/p5_group_classification.csv',index=False)
            pd.DataFrame([self.bridge()]).to_csv(root/'reconciliation/balance_bridge.csv',index=False)
            placeholder=pd.DataFrame({'Описание':['проверка']})
            with ExitStack() as stack:
                stack.enter_context(patch('src.report.read_registry_outputs',return_value={
                    'text':'','rejects':placeholder,'non_subagents':placeholder,'registry':placeholder}))
                stack.enter_context(patch('src.report.summary_tables',return_value=(placeholder,placeholder)))
                for function in ['rejects_table','corporate_table','duplicates_table']:
                    stack.enter_context(patch('src.report.'+function,return_value=placeholder))
                for function in ['read_p5_outputs','read_reconciliation_outputs','subagent_names']:
                    stack.enter_context(patch('src.report.'+function,return_value={}))
                build_report(root,root/'report.xlsx')
            self.assertTrue((root/'p5/accountant_actions.csv').exists())
            w=load_workbook(root/'report.xlsx',read_only=True)
            self.assertIn('Действия бухгалтера',w.sheetnames)
            self.assertIn('Детали действий',w.sheetnames)
            self.assertEqual(w['Действия бухгалтера'].max_row,2)
            w.close()


if __name__=='__main__': unittest.main()
