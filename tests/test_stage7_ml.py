import unittest
import tempfile
import json
from pathlib import Path
from contextlib import ExitStack
from unittest.mock import patch
import numpy as np
import pandas as pd
from src.ml_risk import (LogisticModel, feature_table, make_labels, chronological_split,
                         binary_metrics, aggregate_risks, run_ml, predict_saved_model, report_frames)
from src.schema import ValidationError


class ModelTests(unittest.TestCase):
    def component(self,mid,rid,**changes):
        row=dict(match_id=mid,source_row_id=rid,source='etm',subagent_id='a',date='2026-01-02',
                 amount=100.,original_kind='purchase',creator='bot',ticket10='1234567890')
        row.update(changes); return row

    def labels(self,owners):
        return pd.DataFrame([dict(match_id='m'+str(i),error_owner=o,confidence='medium',assignment_status='hypothesis')
                             for i,o in enumerate(owners)])

    def test_model_learning_and_serialization(self):
        x=np.array([[-2],[-1],[1],[2]],dtype=float); y=np.array([0,0,1,1])
        m=LogisticModel().fit(x,y)
        self.assertGreater(m.predict_proba(x)[-1,1],.7)
        saved=LogisticModel.from_dict(json.loads(json.dumps(m.as_dict())))
        np.testing.assert_allclose(m.predict_proba(x),saved.predict_proba(x))

    def test_scaling_only_uses_training_data(self):
        m=LogisticModel().fit(np.array([[0],[2]]),np.array([0,1]))
        m.predict_proba(np.array([[1000]]))
        self.assertEqual(m.mean[0],1)

    def test_unknown_owner_is_not_negative(self):
        labels=make_labels(self.labels(['none','unclear','bot']))
        self.assertEqual(labels.is_error.iloc[0],0)
        self.assertTrue(pd.isna(labels.is_error.iloc[1]))
        self.assertEqual(labels.is_error.iloc[2],1)

    def test_verified_labels_are_not_mixed_with_rules(self):
        labels=make_labels(self.labels(['none','bot']),pd.DataFrame([
            dict(match_id='m0',is_error=1,error_owner='unclear',is_verified=True)]))
        self.assertEqual(len(labels),1)
        self.assertEqual(labels.is_error.iloc[0],1)
        self.assertEqual(labels.label_basis.iloc[0],'verified')

    def test_inconsistent_manual_label_rejected(self):
        with self.assertRaises(ValidationError):
            make_labels(self.labels(['none']),pd.DataFrame([dict(match_id='m0',is_error=0,error_owner='bot',is_verified=True)]))

    def test_cross_boundary_and_unknown_dates_excluded(self):
        dates=pd.DataFrame({'min':pd.to_datetime(['2026-01-01','2026-02-01','2026-01-31',None]),
                            'max':pd.to_datetime(['2026-01-02','2026-02-02','2026-02-01',None])},index=['a','b','c','d'])
        train,test,excluded,_=chronological_split(dates)
        self.assertEqual(list(train),['a']); self.assertEqual(list(test),['b'])
        self.assertEqual(set(excluded),{'c','d'})

    def test_single_month_rejected(self):
        dates=pd.DataFrame({'min':pd.to_datetime(['2026-01-01']),'max':pd.to_datetime(['2026-01-02'])})
        with self.assertRaises(ValidationError): chronological_split(dates)

    def test_rule_outputs_and_identifiers_not_features(self):
        c=pd.DataFrame([self.component('m1','e1',error_owner='bot',error_type='x',is_error=1,reason='leak')])
        f,_=feature_table(c)
        self.assertFalse(set(f.columns)&{'error_owner','error_type','is_error','reason','match_id','subagent_id','creator'})
        self.assertTrue(np.isfinite(f.to_numpy()).all())

    def test_multiticket_rows_not_counted_twice(self):
        c=pd.DataFrame([self.component('m1','e1',amount=50),self.component('m1','e1',amount=50,ticket10='2345678901')])
        f,_=feature_table(c)
        self.assertAlmostEqual(f.etm_rows.iloc[0],np.log(2))
        self.assertAlmostEqual(f.etm_log_abs_amount.iloc[0],np.log(101))
        p=pd.DataFrame([dict(match_id='m1',p_error=.8,p_source_bot=.7,p_source_agent=.3,p_source_1c=0)])
        transactions,sub,emp=aggregate_risks(p,c)
        self.assertEqual(len(transactions),1)
        self.assertEqual(sub.rows.iloc[0],1)

    def test_metric_ties_and_perfect_prediction(self):
        m=binary_metrics([0,1,0,1],[.5]*4)
        self.assertEqual(m['roc_auc'],.5); self.assertEqual(m['average_precision'],.5)
        m=binary_metrics([0,1],[.1,.9])
        self.assertEqual(m['roc_auc'],1); self.assertEqual(m['macro_f1'],1)
        self.assertIsNone(binary_metrics([0,0],[.2,.3])['roc_auc'])

    def test_transaction_source_probabilities_keep_same_group(self):
        c=pd.DataFrame([self.component('m1','e1'),self.component('m2','e1')])
        p=pd.DataFrame([dict(match_id='m1',p_error=.9,p_source_bot=.8,p_source_agent=.2,p_source_1c=0),
                        dict(match_id='m2',p_error=.7,p_source_bot=.1,p_source_agent=.9,p_source_1c=0)])
        transactions,_,_=aggregate_risks(p,c)
        self.assertEqual(transactions.p_source_bot.iloc[0],.8)
        self.assertEqual(transactions.p_source_agent.iloc[0],.2)
        self.assertAlmostEqual(transactions[['p_source_bot','p_source_agent','p_source_1c']].sum(axis=1).iloc[0],1)

    def test_end_to_end_outputs_and_reloaded_predictions(self):
        with tempfile.TemporaryDirectory() as folder:
            root=Path(folder); (root/'reconciliation').mkdir(); (root/'p5').mkdir()
            rows=[]; labels=[]
            for i in range(40):
                date='2026-01-02' if i<30 else '2026-02-02'
                rows.append(self.component('m'+str(i),'e'+str(i),date=date,amount=100 if i%2 else 200))
                labels.append(dict(match_id='m'+str(i),error_owner='none' if i%2 else 'bot',confidence='medium',assignment_status='hypothesis'))
            c=pd.DataFrame(rows)
            c.to_csv(root/'reconciliation/match_components.csv',index=False)
            pd.DataFrame(labels).to_csv(root/'p5/p5_group_classification.csv',index=False)
            metrics=run_ml(root)
            self.assertEqual(metrics['test_groups'],10)
            self.assertEqual(metrics['source_unsupported_classes'],['1c','agent'])
            saved=predict_saved_model(root/'ml/model.json',c)
            produced=pd.read_csv(root/'ml/group_predictions.csv')
            np.testing.assert_allclose(saved.p_error,produced.p_error)
            self.assertEqual(len(report_frames(root)),4)
            from src.report import build_report
            from openpyxl import load_workbook
            placeholder=pd.DataFrame({'Описание':['проверка']})
            with ExitStack() as stack:
                stack.enter_context(patch('src.report.read_registry_outputs',return_value={
                    'text':'','rejects':placeholder,'non_subagents':placeholder,'registry':placeholder}))
                stack.enter_context(patch('src.report.summary_tables',return_value=(placeholder,placeholder)))
                for function in ['rejects_table','corporate_table','duplicates_table']:
                    stack.enter_context(patch('src.report.'+function,return_value=placeholder))
                for function in ['read_p5_outputs','read_reconciliation_outputs','subagent_names']:
                    stack.enter_context(patch('src.report.'+function,return_value={}))
                build_report(root,root/'report.xlsx')
            w=load_workbook(root/'report.xlsx',read_only=True)
            self.assertIn('Качество модели',w.sheetnames)
            self.assertIn('Приоритет сотрудников',w.sheetnames)
            self.assertEqual(w['Риски транзакций'].max_row,41)
            w.close()


if __name__=='__main__': unittest.main()
