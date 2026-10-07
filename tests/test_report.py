"""Тесты Excel-отчёта (P3).

Как запустить:
    pytest tests/test_report.py -q
"""
import shutil
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

import src.report as report
from src.classify_errors import run_p5
from src.load_registry import load_registry, write_outputs
from src.report import build_report, parse_load_report, row_number
from tests.test_classify_errors import act, etm_op, reg
from tests.test_registry_parser import ETM_KEYS, FIXTURE, TICKET_AGENTS, _write_registry

REPO = Path(__file__).resolve().parent.parent


@pytest.fixture(scope="module")
def interim(tmp_path_factory):
    """Папка interim, собранная загрузчиком из тестовых строк реестра."""
    cases = pd.read_csv(FIXTURE, dtype=str, keep_default_na=False)
    data = tmp_path_factory.mktemp("report_data")
    _write_registry(data, cases)
    result = load_registry(data, etm_keys=ETM_KEYS, ticket_agents=TICKET_AGENTS)
    folder = tmp_path_factory.mktemp("interim")
    write_outputs(result, folder)
    return folder, result


@pytest.fixture(scope="module")
def workbook(interim, tmp_path_factory):
    folder, _ = interim
    path = build_report(folder, tmp_path_factory.mktemp("out") / "report.xlsx")
    return load_workbook(path)


def test_row_number():
    assert row_number("registry:5") == 5


def test_parse_load_report():
    text = "Файл: registry.csv\nИсходных строк: 12\nБез ошибок: 10 (83.33%)\nПричины:\n  ok: 10\n  bad_date: 2\n"
    metrics, reasons = parse_load_report(text)
    assert metrics == [("Файл", "registry.csv"), ("Исходных строк", 12), ("Без ошибок", "10 (83.33%)")]
    assert reasons == [("ok", 10), ("bad_date", 2)]


def test_sheets_in_order(workbook):
    assert workbook.sheetnames == ["Сводка реестра", "Проблемные строки", "Корпоративные", "Дубли реестра"]


def test_summary_has_metrics_and_reasons(workbook, interim):
    _, result = interim
    ws = workbook["Сводка реестра"]
    rows = [[c.value for c in row] for row in ws.iter_rows()]
    assert ["Исходных строк", result.stats["source_rows"]] in [r[:2] for r in rows]
    assert any(r[0] == "ok" and r[1] == "Строка разобрана без ошибок" for r in rows)


def test_rejects_sheet_matches_loader(workbook, interim):
    _, result = interim
    ws = workbook["Проблемные строки"]
    numbers = [row[0].value for row in ws.iter_rows(min_row=2)]
    assert numbers == sorted(row_number(r) for r in result.rejects["row_id"])
    header = [c.value for c in ws[1]]
    assert header[:4] == ["Строка в реестре", "Код ошибки", "Что не так", "Подробности"]


def test_rejects_have_readable_reason(workbook):
    ws = workbook["Проблемные строки"]
    for row in ws.iter_rows(min_row=2):
        assert row[2].value and row[2].value != row[1].value


def test_corporate_sheet_matches_loader(workbook, interim):
    _, result = interim
    ws = workbook["Корпоративные"]
    assert ws.max_row - 1 == len(result.non_subagents)


def test_duplicates_sheet_matches_loader(workbook, interim):
    _, result = interim
    ws = workbook["Дубли реестра"]
    extra = [row[2].value for row in ws.iter_rows(min_row=2)].count("да")
    assert extra == result.stats["dup_extra_rows"]
    assert ws.max_row - 1 == result.stats["dup_rows"]


def test_header_style_and_filter(workbook):
    ws = workbook["Проблемные строки"]
    assert ws["A1"].font.bold
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref.startswith("A1:")


def test_missing_files_fail_loudly(tmp_path):
    with pytest.raises(FileNotFoundError, match="registry.parquet"):
        build_report(tmp_path, tmp_path / "report.xlsx")


# ================================================================
# Листы P5: interim с настоящими результатами P5 на маленьких таблицах
# ================================================================

@pytest.fixture(scope="module")
def full_interim(interim, tmp_path_factory):
    """Копия interim реестра + акты и ETM (для названий) + результаты P5."""
    src, _ = interim
    folder = tmp_path_factory.mktemp("full") / "interim"
    shutil.copytree(src, folder)
    ready = folder / "clean" / "p4_ready"
    ready.mkdir(parents=True)
    acts = pd.DataFrame([
        act("alpha", "2026-01", "Реализация KV AAAAAA", "1000000001", debet=1000.0),
        act("alpha", "2026-01", "Реализация KV BBBBBB", "1000000002", debet=1900.0),   # опечатка 1С
        act("alpha", "2026-01", "Реализация KV CCCCCC", "1000000003", debet=500.0),
        act("alpha", "2026-01", "п/п ЦБ-С000001", credit=5000.0, line_type="payment_bank",
            pay_doc="ЦБ-С000001"),                                                   # нет в ETM
    ]).assign(subagent="ОсОО Альфа Тур")
    etm = pd.DataFrame([
        etm_op("alpha", 1, "purchase", "1000000001", -1000.0),
        etm_op("alpha", 2, "purchase", "1000000002", -1000.0),
        etm_op("alpha", 3, "purchase", "1000000003", -500.0, date="2026-01-10 12:00:00"),
        etm_op("alpha", 4, "purchase", "1000000003", -500.0, date="2026-01-10 15:00:00"),  # дубль бота
    ]).assign(agent="Alpha Tour")
    registry = pd.DataFrame([
        reg("alpha", "r:1", "1000000001", "sale", -1000.0),
        reg("alpha", "r:2", "1000000002", "sale", -1000.0),
        reg("alpha", "r:3", "1000000003", "sale", -500.0),
    ])
    acts.to_csv(ready / "acts_clean.csv", index=False, encoding="utf-8-sig")
    etm.to_csv(ready / "etm_clean.csv", index=False, encoding="utf-8-sig")
    registry_path = folder.parent / "p5_registry.parquet"
    registry.to_parquet(registry_path)
    return folder, run_p5(folder / "clean", registry_path, folder / "p5")


@pytest.fixture(scope="module")
def full_workbook(full_interim, tmp_path_factory):
    folder, _ = full_interim
    return load_workbook(build_report(folder, tmp_path_factory.mktemp("full_out") / "report.xlsx"))


def _column(ws, title, header_row=1):
    header = [c.value for c in ws[header_row]]
    idx = header.index(title)
    return [row[idx].value for row in ws.iter_rows(min_row=header_row + 1)]


def test_full_report_sheets_in_order(full_workbook):
    assert full_workbook.sheetnames == [
        "Сводка", "По субагентам", "Ошибки по типам", "Расхождения по билетам", "Оплаты с ошибкой",
        "Аномалии", "Сотрудники реестра", "Сводка реестра", "Проблемные строки", "Корпоративные",
        "Дубли реестра", "Справочник ошибок",
    ]


def test_overview_counts_match_p5(full_workbook, full_interim):
    _, res = full_interim
    t, p = res["tickets"], res["payments"]
    rows = [[c.value for c in row] for row in full_workbook["Сводка"].iter_rows()]
    assert ["Билетных операций (субагент + билет + продажа/возврат)", len(t)] in [r[:2] for r in rows]
    total = next(r for r in rows if r[0] == "Итого")
    assert total[3] == int(t["is_error"].sum() + p["is_error"].sum())


def test_names_instead_of_keys(full_workbook):
    assert _column(full_workbook["По субагентам"], "Субагент") == ["ОсОО Альфа Тур"]
    assert set(_column(full_workbook["Расхождения по билетам"], "Субагент")) == {"ОсОО Альфа Тур"}


def test_ticket_sheet_has_only_errors(full_workbook, full_interim):
    _, res = full_interim
    t = res["tickets"]
    ws = full_workbook["Расхождения по билетам"]
    assert sorted(_column(ws, "Билет")) == ["1000000002", "1000000003"]
    assert sorted(_column(ws, "Код ошибки")) == sorted(t.loc[t["is_error"] == 1, "error_type"])
    assert set(_column(ws, "Виновник")) == {"1С", "бот ETM"}


def test_payments_and_anomalies_in_russian(full_workbook):
    ws = full_workbook["Оплаты с ошибкой"]
    assert _column(ws, "Что случилось") == ["Оплата есть в 1С, но не зачислена в ETM"]
    assert _column(ws, "Дата в 1С") == ["10.01.2026"]
    ws = full_workbook["Аномалии"]
    assert _column(ws, "Что найдено") == ["Бот повторил операцию (те же билеты, та же сумма)"]
    assert _column(ws, "Важность") == ["высокая"]


def test_money_and_percent_format(full_workbook):
    ws = full_workbook["По субагентам"]
    header = [c.value for c in ws[1]]
    assert ws.cell(row=2, column=header.index("Сумма: 1С, сом") + 1).number_format == "#,##0.00"
    assert ws.cell(row=2, column=header.index("Доля ошибок, %") + 1).number_format == "0.00%"


def test_highlighted_rows(tmp_path):
    path = tmp_path / "h.xlsx"
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        report.write_sheet(writer, "Лист", pd.DataFrame({"a": [1, 2], "b": [3, 4]}), highlight=[False, True])
    ws = load_workbook(path)["Лист"]
    assert ws["A2"].fill.fill_type is None
    assert ws["A3"].fill.fgColor.rgb.endswith("FCE4D6") and ws["B3"].fill.fgColor.rgb.endswith("FCE4D6")


def test_partial_p5_fails_loudly(interim, tmp_path):
    src, _ = interim
    folder = tmp_path / "interim"
    shutil.copytree(src, folder)
    (folder / "p5").mkdir()
    (folder / "p5" / "p5_anomalies.csv").write_text("anomaly_type\n", encoding="utf-8-sig")
    with pytest.raises(FileNotFoundError, match="p5_summary_type.csv"):
        build_report(folder, tmp_path / "report.xlsx")


# ================================================================
# Листы P4: только исправленная версия (с мостом баланса)
# ================================================================

def _write_p4(folder, fixed=True):
    p4 = folder / "p4"
    p4.mkdir(exist_ok=True)
    if not fixed:
        pd.DataFrame({"subagent_id": ["alpha"], "period": ["2026-01"], "bridge_check": [624068.31]}) \
            .to_csv(p4 / "p4_balance_bridge.csv", index=False, encoding="utf-8-sig")
        pd.DataFrame({"subagent_id": ["alpha"], "amount_difference": [-800.0]}) \
            .to_csv(p4 / "p4_subagent_summary.csv", index=False, encoding="utf-8-sig")
        return
    bridge = {"subagent_id": ["alpha", "alpha"], "period": ["2026-01", "2026-02"], "act_exists": [True, True]}
    for col in (["saldo_start", "etm_balance_start", "opening_difference", "closing_difference", "saldo_end",
                 "etm_balance_end"] + list(report.BRIDGE_CAUSES) + report.BRIDGE_ROUNDING):
        bridge[col] = [0.0, 0.0]
    bridge["ops_only_1c"] = [500.0, 0.0]
    bridge["closing_difference"] = [500.0, 1000.0]
    bridge["opening_difference"] = [0.0, 500.0]
    bridge["unexplained"] = [0.0, 500.0]
    pd.DataFrame(bridge).to_csv(p4 / "p4_balance_bridge.csv", index=False, encoding="utf-8-sig")
    summary = {c: [0] for c in ("operations", "matched", "voided", "amount_difference", "only_1c", "only_etm",
                                "period_mismatch", "payments", "payments_only_1c", "payments_only_etm")}
    summary.update(subagent_id=["alpha"], amount_difference_sum=[0.0], only_1c_amount=[500.0],
                   only_etm_amount=[0.0], closing_difference_last=[1000.0], unexplained_total=[500.0],
                   matched_share=[1.0])
    pd.DataFrame(summary).to_csv(p4 / "p4_subagent_summary.csv", index=False, encoding="utf-8-sig")


def test_old_p4_is_not_in_report(full_interim, tmp_path):
    src, _ = full_interim
    folder = tmp_path / "interim"
    shutil.copytree(src, folder)
    _write_p4(folder, fixed=False)
    wb = load_workbook(build_report(folder, tmp_path / "report.xlsx"))
    assert "Мост баланса" not in wb.sheetnames
    assert "Сверка P4 по субагентам" not in wb.sheetnames


def test_fixed_p4_adds_bridge(full_interim, tmp_path):
    src, _ = full_interim
    folder = tmp_path / "interim"
    shutil.copytree(src, folder)
    _write_p4(folder)
    wb = load_workbook(build_report(folder, tmp_path / "report.xlsx"))
    names = wb.sheetnames
    assert names.index("Аномалии") + 1 == names.index("Мост баланса")
    assert names.index("Мост баланса") + 1 == names.index("Сверка P4 по субагентам")
    ws = wb["Мост баланса"]
    assert ws["A1"].value == report.BRIDGE_NOTE
    assert _column(ws, "Субагент", header_row=3) == ["ОсОО Альфа Тур", "ОсОО Альфа Тур"]
    assert _column(ws, "Не объяснено, сом", header_row=3) == [0, 500]
    # Строка с необъяснённой разницей подсвечена, сошедшаяся — нет
    assert ws["A4"].fill.fill_type is None
    assert ws["A5"].fill.fgColor.rgb.endswith("FCE4D6")
    overview = [[c.value for c in row] for row in wb["Сводка"].iter_rows()]
    assert ["из них с необъяснённой разницей", 1] in [r[:2] for r in overview]


# ================================================================
# Реальные данные (если пайплайн уже запускали)
# ================================================================

@pytest.mark.skipif(not (REPO / "interim" / "p5" / "p5_summary_type.csv").exists()
                    or not (REPO / "interim" / "registry_load_report.txt").exists(),
                    reason="нет interim: сначала запустите reconcile.py")
def test_real_data_report(tmp_path):
    p5 = report.read_p5_outputs(REPO / "interim")
    wb = load_workbook(build_report(REPO / "interim", tmp_path / "report.xlsx"), read_only=True)
    assert wb.sheetnames[0] == "Сводка"
    errors = int(p5["tickets"]["is_error"].sum())
    assert wb["Расхождения по билетам"].max_row - 1 == errors
